import re
import struct
from pathlib import Path
import olefile
from pptx import Presentation

from app.services.processors.base import (
    BaseProcessor,
    DocumentProcessingError,
    ExtractionResult,
)


class PPTProcessor(BaseProcessor):
    def extract(self, file_path: Path) -> ExtractionResult:
        if not file_path.exists():
            raise DocumentProcessingError(f"Presentation file not found at {file_path}")

        ext = file_path.suffix.lower()
        if ext == ".pptx":
            return self._extract_pptx(file_path)
        elif ext == ".ppt":
            return self._extract_legacy_ppt(file_path)
        else:
            raise DocumentProcessingError(f"Unsupported presentation extension: {ext}")

    def _extract_pptx(self, file_path: Path) -> ExtractionResult:
        try:
            prs = Presentation(str(file_path))
            slide_count = len(prs.slides)
            sections: list[dict[str, str]] = []

            for idx, slide in enumerate(prs.slides, start=1):
                slide_texts: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            line = "".join(run.text for run in p.runs).strip()
                            if line:
                                slide_texts.append(line)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            row_text = [
                                cell.text.strip() for cell in row.cells if cell.text.strip()
                            ]
                            if row_text:
                                slide_texts.append(" | ".join(row_text))

                # Also check slide notes
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = slide.notes_slide.notes_text_frame.text.strip()
                    if notes_text:
                        slide_texts.append(f"[Notes: {notes_text}]")

                sections.append({
                    "label": f"Slide {idx}",
                    "content": "\n".join(slide_texts).strip(),
                })

            return ExtractionResult(
                sections=sections,
                page_count=None,
                slide_count=slide_count,
                metadata={"format": "PPTX", "slide_count": str(slide_count)},
            )
        except Exception as e:
            raise DocumentProcessingError(f"Failed to process PPTX presentation: {str(e)}", details=str(e))

    def _extract_legacy_ppt(self, file_path: Path) -> ExtractionResult:
        """
        Extract text from legacy PowerPoint 97-2004 binary (.ppt) format.
        Parses OLE compound file and extracts TextCharsAtom (UTF-16LE) and TextBytesAtom (ANSI),
        or performs binary stream string extraction.
        """
        try:
            if not olefile.isOleFile(str(file_path)):
                # If not OLE, fallback to raw binary string extraction
                return self._extract_raw_ppt_text(file_path.read_bytes())

            ole = olefile.OleFileIO(str(file_path))
            try:
                if not ole.exists("PowerPoint Document"):
                    return self._extract_raw_ppt_text(file_path.read_bytes())

                stream = ole.openstream("PowerPoint Document").read()
                return self._parse_ppt_stream(stream)
            finally:
                ole.close()
        except Exception as e:
            raise DocumentProcessingError(f"Failed to process legacy PPT presentation: {str(e)}", details=str(e))

    def _parse_ppt_stream(self, stream: bytes) -> ExtractionResult:
        """
        Parse PowerPoint Document stream for TextCharsAtom (0x0FA0 = 4000)
        and TextBytesAtom (0x0FA8 = 4008).
        """
        texts: list[str] = []
        pos = 0
        length = len(stream)

        while pos + 8 <= length:
            try:
                ver_inst, rec_type, rec_len = struct.unpack("<HHI", stream[pos : pos + 8])
            except Exception:
                break

            pos += 8
            if pos + rec_len > length:
                break

            rec_data = stream[pos : pos + rec_len]

            # 4000 (0x0FA0): RT_TextCharsAtom (UTF-16LE)
            if rec_type == 4000:
                try:
                    txt = rec_data.decode("utf-16-le", errors="ignore").strip()
                    if txt:
                        texts.append(txt)
                except Exception:
                    pass
            # 4008 (0x0FA8): RT_TextBytesAtom (ANSI)
            elif rec_type == 4008:
                try:
                    txt = rec_data.decode("latin-1", errors="ignore").strip()
                    if txt:
                        texts.append(txt)
                except Exception:
                    pass

            pos += rec_len

        if not texts:
            return self._extract_raw_ppt_text(stream)

        sections: list[dict[str, str]] = []
        for idx, t in enumerate(texts, start=1):
            sections.append({"label": f"Slide {idx}", "content": t})

        return ExtractionResult(
            sections=sections,
            page_count=None,
            slide_count=len(sections),
            metadata={"format": "PPT (Legacy)", "slide_count": str(len(sections))},
        )

    def _extract_raw_ppt_text(self, data: bytes) -> ExtractionResult:
        """Fallback extraction for legacy PPT binary streams supporting UTF-16LE and ASCII."""
        text_lines: list[str] = []

        # 1. UTF-16LE extraction
        utf16_matches = re.findall(rb"(?:[\x20-\x7E]\x00){3,}", data)
        for m in utf16_matches:
            try:
                s = m.decode("utf-16-le").strip()
                if s and len(s) > 3 and not s.startswith("PowerPoint") and not s.startswith("Current User"):
                    text_lines.append(s)
            except Exception:
                pass

        # 2. ASCII extraction
        ascii_matches = re.findall(rb"[\x20-\x7E]{3,}", data)
        for m in ascii_matches:
            try:
                s = m.decode("latin-1", errors="ignore").strip()
                if s and len(s) > 3 and not s.startswith("PowerPoint") and not s.startswith("Current User") and s not in text_lines:
                    text_lines.append(s)
            except Exception:
                pass

        content = "\n".join(text_lines)
        return ExtractionResult(
            sections=[{"label": "Slide 1", "content": content}],
            page_count=None,
            slide_count=1,
            metadata={"format": "PPT (Legacy Binary)", "slide_count": "1"},
        )
