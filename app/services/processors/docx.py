import re
import struct
from pathlib import Path
import docx
import olefile

from app.services.processors.base import (
    BaseProcessor,
    DocumentProcessingError,
    ExtractionResult,
)


class DOCXProcessor(BaseProcessor):
    def extract(self, file_path: Path) -> ExtractionResult:
        if not file_path.exists():
            raise DocumentProcessingError(f"Document file not found at {file_path}")

        ext = file_path.suffix.lower()
        if ext == ".docx":
            return self._extract_docx(file_path)
        elif ext == ".doc":
            return self._extract_legacy_doc(file_path)
        else:
            raise DocumentProcessingError(f"Unsupported document extension: {ext}")

    def _extract_docx(self, file_path: Path) -> ExtractionResult:
        try:
            doc = docx.Document(str(file_path))
            sections: list[dict[str, str]] = []
            current_section_label = "Section 1"
            current_lines: list[str] = []
            section_idx = 1

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                # Check if paragraph is a heading
                style_name = p.style.name.lower() if p.style and p.style.name else ""
                if "heading" in style_name or "title" in style_name:
                    if current_lines:
                        sections.append({
                            "label": current_section_label,
                            "content": "\n".join(current_lines).strip(),
                        })
                        section_idx += 1
                        current_lines = []
                    current_section_label = f"Section {section_idx} ({text})"
                    current_lines.append(f"# {text}")
                else:
                    current_lines.append(text)

            # Extract tables as well
            for table in doc.tables:
                table_lines: list[str] = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        table_lines.append(" | ".join(cells))
                if table_lines:
                    current_lines.append("\n[Table]\n" + "\n".join(table_lines))

            if current_lines or not sections:
                sections.append({
                    "label": current_section_label,
                    "content": "\n".join(current_lines).strip(),
                })

            # Read document properties
            doc_meta: dict[str, str] = {"format": "DOCX"}
            try:
                core_props = doc.core_properties
                if core_props.title:
                    doc_meta["title"] = core_props.title
                if core_props.author:
                    doc_meta["author"] = core_props.author
                if core_props.subject:
                    doc_meta["subject"] = core_props.subject
            except Exception:
                pass

            return ExtractionResult(
                sections=sections,
                page_count=None,
                slide_count=None,
                metadata=doc_meta,
            )
        except Exception as e:
            raise DocumentProcessingError(f"Failed to process DOCX document: {str(e)}", details=str(e))

    def _extract_legacy_doc(self, file_path: Path) -> ExtractionResult:
        """
        Extract text from legacy Microsoft Word 97-2004 binary (.doc) format.
        Parses OLE compound structure or extracts printable text from WordDocument stream.
        """
        try:
            if not olefile.isOleFile(str(file_path)):
                return self._extract_raw_doc_text(file_path.read_bytes())

            ole = olefile.OleFileIO(str(file_path))
            try:
                if not ole.exists("WordDocument"):
                    return self._extract_raw_doc_text(file_path.read_bytes())

                stream = ole.openstream("WordDocument").read()
                return self._parse_word_stream(stream)
            finally:
                ole.close()
        except Exception as e:
            raise DocumentProcessingError(f"Failed to process legacy DOC document: {str(e)}", details=str(e))

    def _parse_word_stream(self, stream: bytes) -> ExtractionResult:
        """Extract plain text runs from binary WordDocument stream."""
        if len(stream) < 512:
            return self._extract_raw_doc_text(stream)

        # Word Document FIB (File Information Block)
        # Check FIB magic number: 0xA5EC (Word 97-2004)
        magic = struct.unpack("<H", stream[:2])[0]
        if magic == 0xA5EC:
            # Word 97+ document: text is stored either as 8-bit ANSI or 16-bit Unicode
            # Extract text blocks
            text_blocks: list[str] = []
            
            # Find Unicode text strings (sequences of 2-byte characters where second byte is \x00)
            unicode_matches = re.findall(rb"(?:[\x20-\x7E]\x00){4,}", stream)
            for m in unicode_matches:
                try:
                    s = m.decode("utf-16-le").strip()
                    if s and len(s) > 3:
                        text_blocks.append(s)
                except Exception:
                    pass

            if text_blocks:
                content = "\n\n".join(text_blocks)
                return ExtractionResult(
                    sections=[{"label": "Document Body", "content": content}],
                    page_count=None,
                    slide_count=None,
                    metadata={"format": "DOC (Legacy Binary)"},
                )

        return self._extract_raw_doc_text(stream)

    def _extract_raw_doc_text(self, data: bytes) -> ExtractionResult:
        text_lines: list[str] = []

        # 1. UTF-16LE extraction
        utf16_matches = re.findall(rb"(?:[\x20-\x7E]\x00){3,}", data)
        for m in utf16_matches:
            try:
                s = m.decode("utf-16-le").strip()
                if s and len(s) > 3 and not s.startswith("WordDocument") and not s.startswith("Microsoft"):
                    text_lines.append(s)
            except Exception:
                pass

        # 2. ASCII extraction
        ascii_matches = re.findall(rb"[\x20-\x7E]{4,}", data)
        for m in ascii_matches:
            try:
                s = m.decode("latin-1", errors="ignore").strip()
                if s and len(s) > 3 and not s.startswith("Microsoft") and not s.startswith("WordDocument") and not s.startswith("Normal") and s not in text_lines:
                    text_lines.append(s)
            except Exception:
                pass

        content = "\n".join(text_lines)
        return ExtractionResult(
            sections=[{"label": "Document Body", "content": content}],
            page_count=None,
            slide_count=None,
            metadata={"format": "DOC (Legacy Binary)"},
        )
